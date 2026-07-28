"""Load raw files and construct the standard factor panel.

Data flow
---------
CSV file
    -> inspect_csv()
    -> load_raw_csv()
    -> standardize_long_panel()
    -> panel with MultiIndex(["Date", "Ticker"])

Bloomberg Excel wide-table flow
-------------------------------
Bloomberg Excel wide table
        ↓
Validate file, sheets, widths, headers, and Ticker order
        ↓
Parse 503 price blocks
        ↓
Parse 503 daily descriptor blocks
        ↓
Parse 503 Growth blocks
        ↓
Parse 503 Leverage blocks
        ↓
Normalize Excel serial dates
        ↓
Convert ``#N/A N/A`` and other Bloomberg errors to missing values
        ↓
Record diagnostics
        ↓
Wide table → long table
        ↓
Backward as-of join quarterly descriptors
        ↓
Join GICS industries
        ↓
Standard Date/Ticker panel

Responsibility boundary
-----------------------
This module handles file reading and panel normalization only. It must not
calculate factors, transform exposures, or construct the final matrix X.
"""

from __future__ import annotations

from datetime import date, datetime
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from openpyxl import load_workbook
from openpyxl.utils.datetime import from_excel


STANDARD_INDEX_NAMES = ["Date", "Ticker"]

PRICE_COLUMNS_PER_SECURITY = 2
DESCRIPTOR_COLUMNS_PER_SECURITY = 11

DAILY_DESCRIPTOR_MAPPING = {
    "CUR_MKT_CAP": "market_cap",
    "PX_TO_BOOK_RATIO": "pb",
    "EARN_YLD": "earn_yld",
    "VOLUME": "volume",
    "BETA_ADJ_OVERRIDABLE": "beta",
}

QUARTERLY_DESCRIPTOR_MAPPING = {
    "SALES_GROWTH": "growth_sales",
    "EPS_GROWTH": "growth_eps",
    "TOT_DEBT_TO_COM_EQY": "leverage",
}

INDUSTRY_COLUMN_MAPPING = {
    "gics_sector": "gics_sector",
    "gics_industry_grp": "gics_industry_group",
    "gics_sub_ind": "gics_sub_industry",
}

REQUIRED_US_MARKET_SHEETS = {
    "Config",
    "Universe",
    "Prices",
    "Descriptors_TS",
    "Industries",
    "Meta",
}


def _is_bloomberg_error(value: object) -> bool:
    """Return whether a cached Bloomberg cell contains an error string."""
    return isinstance(value, str) and value.startswith("#")


def _coerce_excel_date(value: object, *, epoch: datetime) -> pd.Timestamp | None:
    """Convert an Excel/cache date value into ``Timestamp``.
    把“看起来像日期、但实际类型不同”的值统一变成 Timestamp
    Empty cells and Bloomberg errors return ``None``. Numeric Excel serials
    are converted using the workbook epoch.
    """
    if value is None or _is_bloomberg_error(value):
        return None

    if isinstance(value, pd.Timestamp):
        return value

    if isinstance(value, datetime):
        return pd.Timestamp(value)

    if isinstance(value, date):
        return pd.Timestamp(value)

    if isinstance(value, (int, float)) and not isinstance(value, bool):
        if pd.isna(value):
            return None
        return pd.Timestamp(from_excel(value, epoch=epoch))

    try:
        return pd.Timestamp(value)
    except (TypeError, ValueError) as exc:
        raise ValueError(f"invalid Excel date value: {value!r}") from exc


def _coerce_numeric(value: object) -> float:
    """Convert a cached numeric value, mapping errors and blanks to NaN."""
    if value is None or _is_bloomberg_error(value):
        return float("nan")

    numeric = pd.to_numeric(value, errors="coerce")
    return float(numeric) if not pd.isna(numeric) else float("nan")


def _require_sheet(workbook: Any, sheet_name: str) -> Any:
    """Return a required worksheet or raise a clear schema error."""
    if sheet_name not in workbook.sheetnames:
        raise ValueError(f"workbook is missing required sheet: {sheet_name}")
    return workbook[sheet_name]


def _read_key_value_sheet(worksheet: Any) -> dict[str, object]:
    """Read a small sheet whose first two columns contain key/value pairs."""
    result: dict[str, object] = {}
    for row in worksheet.iter_rows(values_only=True):
        if row and row[0] is not None:
            result[str(row[0])] = row[1] if len(row) > 1 else None
    return result


def _validate_csv_path(path: str | Path) -> Path:
    """Validate a CSV path and return an absolute ``Path``.
    Acceptance criteria
    -------------------
    - Both ``str`` and ``Path`` inputs work.
    - Missing paths, directories, and non-CSV files are rejected.
    - The returned value is an absolute ``Path``.
    """
    # 1. Convert ``path`` into a ``Path`` object.
    path = Path(path)

    # 2. Expand ``~`` and resolve it into an absolute path.
    path = path.expanduser().resolve()

    # 3. Raise ``FileNotFoundError`` if the path does not exist.
    if not path.exists():
        raise FileNotFoundError(f"CSV file not found: {path}")

    # 4. Raise ``ValueError`` if the path is not a regular file.
    if not path.is_file():
        raise ValueError(f"CSV path is not a regular file: {path}")

    # 5. Raise ``ValueError`` unless its suffix is ``.csv``.
    if not path.suffix.casefold() == ".csv":
        raise ValueError(f"CSV path must have a .csv suffix: {path}")
    
    return path
    


def inspect_csv(
    path: str | Path,
    *,
    nrows: int = 10,
    encoding: str = "utf-8",
    **read_csv_kwargs: Any,
) -> pd.DataFrame:
    """Read a small, unmodified preview of a CSV.

    Questions to investigate
    ------------------------
    - What does one giant column suggest about the delimiter?
    - What should change if the header appears as a data row?
    - What is the difference between ``header`` and ``skiprows``?

    Acceptance criteria
    -------------------
    - Large CSV files can be inspected without loading the complete file.
    - Invalid ``nrows`` values are rejected.
    - Options such as ``sep=";"`` reach ``pd.read_csv``.
    """
    # 1. Confirm ``nrows`` is a positive integer.Investigate why ``bool`` needs special consideration.
    if (
        not isinstance(nrows, int)
        or isinstance(nrows, bool)
        or nrows <= 0
    ):
        raise ValueError(
            f"nrows must be a positive integer, got {nrows!r}"
        )
        
    # 2. Validate ``path`` with ``_validate_csv_path``.
    path = _validate_csv_path(path)
    # 3. Call ``pd.read_csv`` with the validated path.
    df = pd.read_csv(path, nrows=nrows, encoding=encoding, **read_csv_kwargs)
 
    # 6. Return the preview without renaming, parsing project-specific dates,
    #    constructing an index, or filling NaN.
    return df

def describe_csv(preview: pd.DataFrame) -> pd.DataFrame:
    """Return a column-level summary of a CSV preview.

    Expected output columns
    -----------------------
    - ``column``
    - ``inferred_dtype``
    - ``non_null_preview_rows``
    - ``null_preview_rows``

    Acceptance criteria
    -------------------
    - The summary has one row per preview column.
    - Summary order matches the CSV column order.
    - Null count plus non-null count equals ``len(preview)``.
    """ 

    #1. Confirm ``preview`` is a DataFrame.
    if not isinstance(preview, pd.DataFrame):
        raise TypeError(f"preview must be a DataFrame: {type(preview).__name__}")
    #2. Collect column names in their original order.
    columns = [str(column) for column in preview.columns]
    #3. Collect the dtype inferred for each column.
    dtypes = [str(dtype) for dtype in preview.dtypes]
    #4. Count null and non-null preview values per column.
    non_null_counts = (preview.count().tolist())
    null_counts = [(len(preview) - count) for count in non_null_counts]
    
    #5. Assemble and return a new summary DataFrame.
    return pd.DataFrame({
        "column": columns,
        "inferred_dtype": dtypes,
        "non_null_preview_rows": non_null_counts,
        "null_preview_rows": null_counts,
    })




def load_raw_csv(
    path: str | Path,
    *,
    encoding: str = "utf-8",
    **read_csv_kwargs: Any,
) -> pd.DataFrame:
    """Read the complete CSV into an initial, unmodified DataFrame.
    Acceptance criteria
    -------------------
    - Row order and raw column names are preserved.
    - Missing values remain missing.
    - CSV parsing options can be supplied by the caller.
    """
    # 1.Validate ``path`` with ``_validate_csv_path``.
    path = _validate_csv_path(path)
    # 2. Read the complete file with ``pd.read_csv``.
    df = pd.read_csv(path, encoding=encoding, **read_csv_kwargs)
    # 3. Return the raw DataFrame.
    return df


def standardize_long_panel(
    raw: pd.DataFrame,
    *,
    date_column: str,
    ticker_column: str,
    column_mapping: dict[str, str] | None = None,
) -> pd.DataFrame:
    """Convert a long-form raw table into the standard factor panel.

    Example mapping
    ---------------
    ``{"CUR_MKT_CAP": "market_cap", "PX_LAST": "px_last"}``

    Acceptance criteria
    -------------------
    - ``panel.index.names == STANDARD_INDEX_NAMES``.
    - Date is datetime-like.
    - Every ``(Date, Ticker)`` pair is unique.
    - The index is sorted.
    - Feature NaN values are preserved.
    - ``raw`` is not modified.
    - A numeric ``market_cap`` result can be passed to ``SizeFactor``.
    """
    # 1. Confirm ``raw`` is a DataFrame.
    if not isinstance(raw, pd.DataFrame):
        raise TypeError(f"raw must be a DataFrame: {type(raw).__name__}")
    # 2. Confirm ``date_column`` and ``ticker_column`` are non-empty strings.
    if not isinstance(date_column, str) or not date_column.strip():
        raise ValueError(f"date_column must be a non-empty string: {date_column}")
    if not isinstance(ticker_column, str) or not ticker_column.strip():
        raise ValueError(f"ticker_column must be a non-empty string: {ticker_column}")
    if date_column == ticker_column:
        raise ValueError("date_column and ticker_column must be different")
    # 3. Confirm both identifier columns exist.
    if date_column not in raw.columns:
        raise ValueError(f"date_column not found in raw columns: {date_column}")
    if ticker_column not in raw.columns:
        raise ValueError(f"ticker_column not found in raw columns: {ticker_column}")
    # 4. Treat ``column_mapping=None`` as an empty mapping.
    if column_mapping is None:
        column_mapping = {}
    elif not isinstance(column_mapping, dict):
        raise TypeError(
            "column_mapping must be a dict or None, "
            f"got {type(column_mapping).__name__}"
        )
    # 5. Confirm all source columns in ``column_mapping`` exist.
    missing_source_columns = sorted(
        set(column_mapping) - set(raw.columns)
    )
    if missing_source_columns:
        raise ValueError(
            "column_mapping contains source columns not found in raw: "
            f"{missing_source_columns}"
        )

    mapped_identifiers = sorted(
        {date_column, ticker_column}.intersection(column_mapping)
    )
    if mapped_identifiers:
        raise ValueError(
            "column_mapping must not rename identifier columns: "
            f"{mapped_identifiers}"
        )

    # 6. Confirm renaming will not create duplicate column names.
    rename_mapping = {
        **column_mapping,
        date_column: "Date",
        ticker_column: "Ticker",
    }
    renamed_columns = [
        rename_mapping.get(column, column)
        for column in raw.columns
    ]
    if len(renamed_columns) != len(set(renamed_columns)):
        raise ValueError(
            "column renaming would create duplicate output column names: "
            f"{renamed_columns}"
        )

    # 7. Work on a deep copy so ``raw`` remains unchanged.
    panel = raw.copy(deep=True)
    # 8-9. Apply feature and identifier mappings together.
    panel = panel.rename(columns=rename_mapping)
    # 10. Convert ``Date`` with ``pd.to_datetime`` and expose invalid dates
    panel["Date"] = pd.to_datetime(panel["Date"], errors="raise")
    # 11. Strip surrounding whitespace from Ticker values.
    non_null_tickers = panel["Ticker"].dropna()
    if not non_null_tickers.map(
        lambda ticker: isinstance(ticker, str)
    ).all():
        raise TypeError("all non-null Ticker values must be strings")
    panel["Ticker"] = panel["Ticker"].str.strip()
    # 12. Do not invent ``"US Equity"`` unless the raw schema proves only short
    #     Bloomberg symbols were supplied.
    # 13. Reject missing or empty Date/Ticker identifiers.
    invalid_identifiers = (
        panel["Date"].isna()
        | panel["Ticker"].isna()
        | panel["Ticker"].eq("")
    )
    if invalid_identifiers.any():
        invalid_count = int(invalid_identifiers.sum())
        raise ValueError(
            "Date and Ticker must not be missing or empty; "
            f"found {invalid_count} invalid rows"
        )
    # 14. Set ``MultiIndex(["Date", "Ticker"])``.
    panel = panel.set_index(["Date", "Ticker"])
    # 15. Reject duplicated ``(Date, Ticker)`` observations.
    if panel.index.has_duplicates:
        raise ValueError(f"panel index must contain unique (Date, Ticker) pairs")
    # 16. Sort by Date and then Ticker.
    panel = panel.sort_index()
    # 17. Return the panel without filling feature NaN values.
    return panel


def _validate_excel_path(path: str | Path) -> Path:
    """Validate a Bloomberg Excel path and return an absolute ``Path``.

    Implementation outline
    ----------------------
    1. Convert ``path`` into ``Path``.
    2. Expand ``~`` and resolve it.
    3. Reject missing paths.
    4. Reject directories.
    5. Accept only ``.xlsx``; compare the suffix case-insensitively.
    6. Return the absolute path.

    Acceptance criteria
    -------------------
    - Both string and Path inputs work.
    - Missing files, directories, CSV files, and legacy ``.xls`` files fail
      with clear exceptions.
    """
    excel_path = Path(path).expanduser().resolve()

    if not excel_path.exists():
        raise FileNotFoundError(f"Excel file does not exist: {excel_path}")

    if not excel_path.is_file():
        raise ValueError(f"Excel path must point to a file: {excel_path}")

    if excel_path.suffix.casefold() != ".xlsx":
        raise ValueError(
            f"expected an .xlsx file, got "
            f"{excel_path.suffix or 'no extension'}: {excel_path}"
        )

    return excel_path


def inspect_us_market_workbook(path: str | Path) -> dict[str, object]:
    """Inspect the Bloomberg workbook without parsing the complete dataset.

    Expected output
    ---------------
    Return a dictionary containing at least:

    - ``sheet_names``
    - ``sheet_dimensions``
    - ``config``
    - ``meta``

    Implementation outline
    ----------------------
    1. Validate the path.
    2. Open it with ``openpyxl.load_workbook`` in read-only mode.
    3. Use ``data_only=True`` because factor construction needs Bloomberg's
       cached values, not the ``BDH``/``BDP`` formula strings.
    4. Confirm all required sheets exist:
       ``Config``, ``Universe``, ``Prices``, ``Descriptors_TS``,
       ``Industries``, ``Meta``.
    5. Record each sheet's dimensions.
    6. Read the small Config and Meta sheets into dictionaries.
    7. Close the workbook even if validation fails.
    8. Return the inspection dictionary.

    Questions to investigate
    ------------------------
    - Why can openpyxl read cached Bloomberg values but not calculate BDH?
    - What happens if the workbook was saved before Bloomberg finished
      refreshing?
    - Why should inspection use read-only mode?

    Acceptance criteria
    -------------------
    - The function never modifies the workbook.
    - Missing required sheets fail clearly.
    - Config reports SPX Index, daily period, and the expected date range.
    - Meta reports 503 securities, two price columns per security, and eleven
      descriptor columns per security.
    """
    excel_path = _validate_excel_path(path)
    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
        keep_links=True,
    )
    try:
        missing_sheets = sorted(
            REQUIRED_US_MARKET_SHEETS - set(workbook.sheetnames)
        )
        if missing_sheets:
            raise ValueError(
                f"workbook is missing required sheets: {missing_sheets}"
            )

        dimensions = {
            worksheet.title: worksheet.calculate_dimension()
            for worksheet in workbook.worksheets
        }
        config = _read_key_value_sheet(workbook["Config"])
        meta = _read_key_value_sheet(workbook["Meta"])

        return {
            "path": excel_path,
            "sheet_names": list(workbook.sheetnames),
            "sheet_dimensions": dimensions,
            "config": config,
            "meta": meta,
        }
    finally:
        workbook.close()


def load_universe_sheet(workbook: Any) -> list[str]:
    """Read and validate the ordered Bloomberg Ticker universe.

    Implementation outline
    ----------------------
    1. Read ``Universe`` using cached cell values.
    2. Confirm cell A1 is ``ticker``.
    3. Read non-empty values from A2 downward.
    4. Confirm every Ticker is a non-empty string.
    5. Strip surrounding whitespace.
    6. Reject duplicated Tickers.
    7. Confirm the count agrees with ``Meta.n_securities``.
    8. Preserve the workbook order and return ``list[str]``.

    Acceptance criteria
    -------------------
    - The supplied workbook returns exactly 503 unique Tickers.
    - ``A US Equity`` and ``AAPL US Equity`` remain distinct identifiers.
    - No Ticker is invented or silently removed.
    """
    worksheet = _require_sheet(workbook, "Universe")
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if not header or header[0] != "ticker":
        raise ValueError(
            f"Universe!A1 must be 'ticker', got "
            f"{None if not header else header[0]!r}"
        )

    tickers: list[str] = []
    for row_number, row in enumerate(rows, start=2):
        value = row[0] if row else None
        if value is None:
            continue
        if not isinstance(value, str) or not value.strip():
            raise ValueError(
                f"Universe!A{row_number} must contain a non-empty Ticker"
            )
        tickers.append(value.strip())

    duplicates = sorted(
        ticker for ticker in set(tickers) if tickers.count(ticker) > 1
    )
    if duplicates:
        raise ValueError(f"Universe contains duplicate Tickers: {duplicates}")

    meta = _read_key_value_sheet(_require_sheet(workbook, "Meta"))
    expected_count = meta.get("n_securities")
    if expected_count is not None and len(tickers) != int(expected_count):
        raise ValueError(
            f"Universe contains {len(tickers)} Tickers, "
            f"but Meta expects {expected_count}"
        )

    return tickers


def load_prices_sheet(
    workbook: Any,
    tickers: list[str],
) -> pd.DataFrame:
    """Convert the wide Prices sheet into a long price panel.

    Expected output
    ---------------
    A DataFrame with:

    - index: ``MultiIndex(["Date", "Ticker"])``
    - columns: ``["px_last"]``

    Acceptance criteria
    -------------------
    - The full-history securities have 502 dated observations.
    - Partial-history securities remain partial; they are not backfilled.
    - ``FDXF US Equity`` and ``HONA US Equity`` do not crash parsing.
    - Bloomberg error strings never enter ``px_last`` as text.
    - The returned ``px_last`` column is numeric.
    """
    # 1. Read Prices with openpyxl cached values (workbook was opened with
    # data_only=True by the public loader).
    worksheet = _require_sheet(workbook, "Prices")

    # 2. Confirm there are exactly two columns per security.
    expected_width = len(tickers) * PRICE_COLUMNS_PER_SECURITY
    if worksheet.max_column != expected_width:
        raise ValueError(
            f"Prices must contain {expected_width} columns, "
            f"got {worksheet.max_column}"
        )

    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError("Prices sheet is empty")

    # 3. For security i, its [Date, PX_LAST] block begins at column 2*i.
    # 4. Confirm the Ticker stored in row 1 matches Universe[i].
    for security_index, ticker in enumerate(tickers):
        ticker_header = header[2 * security_index + 1]
        if ticker_header != ticker:
            raise ValueError(
                f"Prices block {security_index} expected Ticker "
                f"{ticker!r}, got {ticker_header!r}"
            )

    dates: list[pd.Timestamp] = []
    output_tickers: list[str] = []
    prices: list[float] = []
    diagnostics: list[dict[str, object]] = []

    # 5. Read the data rows (row 2 onward) from every two-column block.
    for row_number, row in enumerate(rows, start=2):
        for security_index, ticker in enumerate(tickers):
            date_value = row[2 * security_index]
            price_value = row[2 * security_index + 1]

            # 6. Record Bloomberg errors; conversion below turns them into
            # missing values instead of keeping '#N/A N/A' as a string.
            if _is_bloomberg_error(date_value):
                diagnostics.append(
                    {
                        "sheet": "Prices",
                        "row": row_number,
                        "ticker": ticker,
                        "field": "Date",
                        "error": date_value,
                    }
                )
            if _is_bloomberg_error(price_value):
                diagnostics.append(
                    {
                        "sheet": "Prices",
                        "row": row_number,
                        "ticker": ticker,
                        "field": "PX_LAST",
                        "error": price_value,
                    }
                )

            # 7. Keep only rows that actually contain a valid Date.
            # 8. Normalize datetime / Excel serial dates.
            parsed_date = _coerce_excel_date(
                date_value,
                epoch=workbook.epoch,
            )
            if parsed_date is None:
                continue

            # 9. Add this block's Ticker to the long-form output row.
            dates.append(parsed_date)
            output_tickers.append(ticker)

            # 8. Convert PX_LAST into a numeric value; blanks/errors become NaN.
            prices.append(_coerce_numeric(price_value))

    # 10. Concatenate every security's two-column block into one long table.
    # 12. Construct the standard Date/Ticker MultiIndex.
    panel = pd.DataFrame(
        {
            "Date": dates,
            "Ticker": output_tickers,
            "px_last": prices,
        }
    ).set_index(STANDARD_INDEX_NAMES)

    # 11. Reject duplicated observations before handing prices downstream.
    if panel.index.has_duplicates:
        raise ValueError("Prices contains duplicate (Date, Ticker) pairs")

    # 12. Sort the final Date/Ticker index.
    panel = panel.sort_index()
    panel.attrs["diagnostics"] = diagnostics
    return panel


def load_descriptors_sheet(
    workbook: Any,
    tickers: list[str],
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Parse daily and quarterly descriptor blocks.

    Expected output
    ---------------
    Return ``(daily_descriptors, quarterly_descriptors)``.

    ``daily_descriptors``:

    - index: ``MultiIndex(["Date", "Ticker"])``
    - columns: ``market_cap``, ``pb``, ``earn_yld``, ``volume``, ``beta``

    ``quarterly_descriptors``:

    - index: ``MultiIndex(["Date", "Ticker"])``
    - columns: ``growth_sales``, ``growth_eps``, ``leverage``

    Point-in-time warning
    ---------------------
    The quarterly Date values are Bloomberg ``Period="Q"`` observation dates.
    Do not label them as filing dates or confirmed public-availability dates.
    Record their semantics as ``PIT unverified`` until independently checked.

    Acceptance criteria
    -------------------
    - Daily fields cover 2024-01-02 through 2025-12-31 where available.
    - Growth has up to eight quarterly observations per security.
    - Excel serial ``45381`` is parsed as a date, not kept as an integer.
    - ``#N/A N/A`` becomes NaN, including leverage values.
    - Output columns use AlphaStream names from the mapping constants.
    """
    worksheet = _require_sheet(workbook, "Descriptors_TS")

    # 1. Confirm the sheet width is 11 columns per security.
    expected_width = len(tickers) * DESCRIPTOR_COLUMNS_PER_SECURITY
    if worksheet.max_column != expected_width:
        raise ValueError(
            f"Descriptors_TS must contain {expected_width} columns, "
            f"got {worksheet.max_column}"
        )

    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)
    if header is None:
        raise ValueError("Descriptors_TS sheet is empty")

    # 3. Define the five raw fields in each daily descriptor sub-block.
    daily_raw_fields = list(DAILY_DESCRIPTOR_MAPPING)

    # 4. Define the two raw fields in each quarterly Growth sub-block.
    growth_raw_fields = [
        "SALES_GROWTH",
        "EPS_GROWTH",
    ]
    # 5. Define the single raw field in each quarterly Leverage sub-block.
    leverage_raw_fields = ["TOT_DEBT_TO_COM_EQY"]

    # 2. For security i, its eleven-column block starts at i * 11.
    for security_index, ticker in enumerate(tickers):
        base = security_index * DESCRIPTOR_COLUMNS_PER_SECURITY

        # 6. Validate each daily/Growth/Leverage sub-block Ticker header.
        # 7. Validate raw descriptor names against the mapping constants.
        checks = [
            (base, ticker),
            *[
                (base + 1 + offset, field)
                for offset, field in enumerate(daily_raw_fields)
            ],
            (base + 6, ticker),
            *[
                (base + 7 + offset, field)
                for offset, field in enumerate(growth_raw_fields)
            ],
            (base + 9, ticker),
            (base + 10, leverage_raw_fields[0]),
        ]
        for column_index, expected_value in checks:
            if header[column_index] != expected_value:
                raise ValueError(
                    f"Descriptors_TS column {column_index + 1} expected "
                    f"{expected_value!r}, got {header[column_index]!r}"
                )

    daily_records: list[dict[str, object]] = []
    growth_records: list[dict[str, object]] = []
    leverage_records: list[dict[str, object]] = []
    diagnostics: list[dict[str, object]] = []

    for row_number, row in enumerate(rows, start=2):
        for security_index, ticker in enumerate(tickers):
            base = security_index * DESCRIPTOR_COLUMNS_PER_SECURITY

            # 3. Parse the daily [Date, 5 descriptors] sub-block.
            daily_date_raw = row[base]

            # 10-11. Normal datetime and Excel serial dates are normalized
            # through _coerce_excel_date (e.g. AAPL leverage serial 45381).
            daily_date = _coerce_excel_date(
                daily_date_raw,
                epoch=workbook.epoch,
            )
            # 8. Record Bloomberg errors; _coerce_excel_date maps them to None.
            if _is_bloomberg_error(daily_date_raw):
                diagnostics.append(
                    {
                        "sheet": "Descriptors_TS",
                        "row": row_number,
                        "ticker": ticker,
                        "field": "daily_date",
                        "error": daily_date_raw,
                    }
                )
            if daily_date is not None:
                record: dict[str, object] = {
                    "Date": daily_date,
                    "Ticker": ticker,
                }
                for offset, raw_field in enumerate(daily_raw_fields, start=1):
                    raw_value = row[base + offset]
                    # 8. Record Bloomberg field errors for diagnostics.
                    if _is_bloomberg_error(raw_value):
                        diagnostics.append(
                            {
                                "sheet": "Descriptors_TS",
                                "row": row_number,
                                "ticker": ticker,
                                "field": raw_field,
                                "error": raw_value,
                            }
                        )
                    # 9 & 12. Convert numeric values; blanks/errors become NaN.
                    record[DAILY_DESCRIPTOR_MAPPING[raw_field]] = (
                        _coerce_numeric(raw_value)
                    )
                daily_records.append(record)

            # 4. Parse the quarterly Growth [Date, SALES_GROWTH, EPS_GROWTH]
            # sub-block.
            growth_date_raw = row[base + 6]

            # 10-11. Normalize datetime cells and any Excel serial dates.
            growth_date = _coerce_excel_date(
                growth_date_raw,
                epoch=workbook.epoch,
            )
            # 8. Record a Bloomberg error date before it becomes missing.
            if _is_bloomberg_error(growth_date_raw):
                diagnostics.append(
                    {
                        "sheet": "Descriptors_TS",
                        "row": row_number,
                        "ticker": ticker,
                        "field": "growth_date",
                        "error": growth_date_raw,
                    }
                )
            if growth_date is not None:
                growth_record: dict[str, object] = {
                    "Date": growth_date,
                    "Ticker": ticker,
                }
                for offset, raw_field in enumerate(
                    growth_raw_fields,
                    start=7,
                ):
                    raw_value = row[base + offset]
                    # 8. Record Bloomberg field errors for diagnostics.
                    if _is_bloomberg_error(raw_value):
                        diagnostics.append(
                            {
                                "sheet": "Descriptors_TS",
                                "row": row_number,
                                "ticker": ticker,
                                "field": raw_field,
                                "error": raw_value,
                            }
                        )
                    # 9 & 12. Convert numeric values; blanks/errors become NaN.
                    growth_record[
                        QUARTERLY_DESCRIPTOR_MAPPING[raw_field]
                    ] = _coerce_numeric(raw_value)
                growth_records.append(growth_record)

            # 5. Parse the quarterly Leverage [Date, TOT_DEBT_TO_COM_EQY]
            # sub-block.
            leverage_date_raw = row[base + 9]

            # 10-11. Normalize datetime cells and Excel serial dates. AAPL's
            # first leverage date is serial 45381 in the current workbook.
            leverage_date = _coerce_excel_date(
                leverage_date_raw,
                epoch=workbook.epoch,
            )
            leverage_value_raw = row[base + 10]
            # 8. Record a Bloomberg error date before it becomes missing.
            if _is_bloomberg_error(leverage_date_raw):
                diagnostics.append(
                    {
                        "sheet": "Descriptors_TS",
                        "row": row_number,
                        "ticker": ticker,
                        "field": "leverage_date",
                        "error": leverage_date_raw,
                    }
                )
            if leverage_date is not None:
                # 8. Record Bloomberg field errors for diagnostics.
                if _is_bloomberg_error(leverage_value_raw):
                    diagnostics.append(
                        {
                            "sheet": "Descriptors_TS",
                            "row": row_number,
                            "ticker": ticker,
                            "field": leverage_raw_fields[0],
                            "error": leverage_value_raw,
                        }
                    )
                leverage_records.append(
                    {
                        "Date": leverage_date,
                        "Ticker": ticker,
                        # 9 & 12. Convert numeric values; blanks/errors become NaN.
                        "leverage": _coerce_numeric(leverage_value_raw),
                    }
                )

    # 13. Concatenate all per-security blocks into sorted Date/Ticker tables.
    daily = pd.DataFrame.from_records(daily_records)
    if daily.empty:
        daily = pd.DataFrame(
            columns=[
                *STANDARD_INDEX_NAMES,
                *DAILY_DESCRIPTOR_MAPPING.values(),
            ]
        )
    daily = daily.set_index(STANDARD_INDEX_NAMES).sort_index()
    # 14. Reject duplicate observations in every output table.
    if daily.index.has_duplicates:
        raise ValueError(
            "daily descriptors contain duplicate (Date, Ticker) pairs"
        )

    growth = pd.DataFrame.from_records(growth_records)
    if growth.empty:
        growth = pd.DataFrame(
            columns=[
                *STANDARD_INDEX_NAMES,
                "growth_sales",
                "growth_eps",
            ]
        )
    growth = growth.set_index(STANDARD_INDEX_NAMES).sort_index()

    leverage = pd.DataFrame.from_records(leverage_records)
    if leverage.empty:
        leverage = pd.DataFrame(
            columns=[
                *STANDARD_INDEX_NAMES,
                "leverage",
            ]
        )
    leverage = leverage.set_index(STANDARD_INDEX_NAMES).sort_index()

    # 14. Apply the same duplicate check to Growth and Leverage tables.
    if growth.index.has_duplicates:
        raise ValueError(
            "growth descriptors contain duplicate (Date, Ticker) pairs"
        )
    if leverage.index.has_duplicates:
        raise ValueError(
            "leverage descriptors contain duplicate (Date, Ticker) pairs"
        )

    quarterly = growth.join(leverage, how="outer").sort_index()
    quarterly.attrs["point_in_time_status"] = "unverified"
    quarterly.attrs["date_semantics"] = (
        "Bloomberg Period=Q observation date; not confirmed filing date"
    )
    daily.attrs["diagnostics"] = diagnostics
    quarterly.attrs["diagnostics"] = diagnostics
    return daily, quarterly


def load_industries_sheet(
    workbook: Any,
    tickers: list[str],
) -> pd.DataFrame:
    """Load the static GICS mapping.

    Expected output
    ---------------
    A DataFrame indexed by ``Ticker`` with columns:

    - ``gics_sector``
    - ``gics_industry_group``
    - ``gics_sub_industry``

    Acceptance criteria
    -------------------
    - The current workbook returns 503 unique rows.
    - No GICS field is missing.
    - Truncated cached Bloomberg labels are preserved rather than guessed.
    """
    # 1. Read Industries with cached Bloomberg values.
    worksheet = _require_sheet(workbook, "Industries")
    rows = worksheet.iter_rows(values_only=True)
    header = next(rows, None)

    # 2. Validate the raw four-column Bloomberg header.
    expected_header = [
        "ticker",
        "gics_sector",
        "gics_industry_grp",
        "gics_sub_ind",
    ]
    if header is None or list(header[:4]) != expected_header:
        raise ValueError(
            f"Industries header must be {expected_header}, "
            f"got {None if header is None else list(header[:4])}"
        )

    records: list[dict[str, object]] = []
    for row_number, row in enumerate(rows, start=2):
        ticker = row[0]

        # 4. Validate Ticker strings and strip surrounding whitespace.
        # 5. Reject rows with missing or invalid Tickers.
        if ticker is None:
            raise ValueError(f"Industries!A{row_number} is missing Ticker")
        if not isinstance(ticker, str) or not ticker.strip():
            raise ValueError(
                f"Industries!A{row_number} must contain a string Ticker"
            )

        record: dict[str, object] = {"Ticker": ticker.strip()}
        for column_index, raw_name in enumerate(expected_header[1:], start=1):
            value = row[column_index]
            if value is None or _is_bloomberg_error(value):
                raise ValueError(
                    f"Industries row {row_number} has invalid "
                    f"{raw_name}: {value!r}"
                )
            # 3. Rename raw GICS fields using the common mapping constant.
            record[INDUSTRY_COLUMN_MAPPING[raw_name]] = value
        records.append(record)

    # 8. Set Ticker as the table index.
    industries = pd.DataFrame.from_records(records).set_index("Ticker")

    # 5. Reject duplicate Tickers after all rows have been collected.
    if industries.index.has_duplicates:
        raise ValueError("Industries contains duplicate Tickers")

    # 6. Confirm Industries has the same Ticker set as Universe.
    # 7. The equality check also preserves and validates Universe order.
    if list(industries.index) != tickers:
        missing = sorted(set(tickers) - set(industries.index))
        extra = sorted(set(industries.index) - set(tickers))
        raise ValueError(
            "Industries Tickers/order do not match Universe; "
            f"missing={missing}, extra={extra}"
        )

    # 8. Return the indexed, Universe-ordered GICS table.
    return industries


def build_us_market_panel(
    prices: pd.DataFrame,
    daily_descriptors: pd.DataFrame,
    quarterly_descriptors: pd.DataFrame,
    industries: pd.DataFrame,
) -> pd.DataFrame:
    """Combine parsed tables into the standard factor input panel.

    Important
    ---------
    A backward as-of join prevents using a later observation date, but it does
    not prove that the quarterly observation was publicly available on that
    date. Keep the point-in-time warning in project metadata.

    Acceptance criteria
    -------------------
    - Index names are exactly ``["Date", "Ticker"]``.
    - Price rows are not created or removed during feature joins.
    - Quarterly values never flow backward before their observation date.
    - Industries are constant within Ticker.
    - Missing factor inputs remain NaN.
    """
    # 1. Define and validate each input's required index and columns.
    expected_price_columns = {"px_last"}
    expected_daily_columns = set(DAILY_DESCRIPTOR_MAPPING.values())
    expected_quarterly_columns = set(
        QUARTERLY_DESCRIPTOR_MAPPING.values()
    )
    expected_industry_columns = set(INDUSTRY_COLUMN_MAPPING.values())

    for name, table, columns in [
        ("prices", prices, expected_price_columns),
        ("daily_descriptors", daily_descriptors, expected_daily_columns),
        (
            "quarterly_descriptors",
            quarterly_descriptors,
            expected_quarterly_columns,
        ),
    ]:
        if not isinstance(table, pd.DataFrame):
            raise TypeError(f"{name} must be a pandas DataFrame")
        if not isinstance(table.index, pd.MultiIndex):
            raise ValueError(f"{name} must use a MultiIndex")
        if list(table.index.names) != STANDARD_INDEX_NAMES:
            raise ValueError(
                f"{name} index names must be {STANDARD_INDEX_NAMES}"
            )
        if table.index.has_duplicates:
            raise ValueError(f"{name} index must be unique")
        missing_columns = sorted(columns - set(table.columns))
        if missing_columns:
            raise ValueError(
                f"{name} is missing columns: {missing_columns}"
            )

    if not isinstance(industries, pd.DataFrame):
        raise TypeError("industries must be a pandas DataFrame")
    if industries.index.name != "Ticker":
        raise ValueError("industries index must be named 'Ticker'")
    if industries.index.has_duplicates:
        raise ValueError("industries index must be unique")
    missing_industry_columns = sorted(
        expected_industry_columns - set(industries.columns)
    )
    if missing_industry_columns:
        raise ValueError(
            f"industries is missing columns: {missing_industry_columns}"
        )

    # 2. Use Prices as the authoritative Date/Ticker observation grid.
    original_index = prices.index

    prices_for_join = prices.copy(deep=False)
    prices_for_join.attrs = {}
    daily_for_join = daily_descriptors.copy(deep=False)
    daily_for_join.attrs = {}
    quarterly_for_join = quarterly_descriptors.copy(deep=False)
    quarterly_for_join.attrs = {}
    industries_for_join = industries.copy(deep=False)
    industries_for_join.attrs = {}

    # 3. Join daily descriptors only when Date and Ticker match exactly.
    panel = prices_for_join.join(daily_for_join, how="left")

    panel_rows = panel.reset_index()
    panel_rows.attrs = {}
    quarterly_rows = quarterly_for_join.reset_index()
    quarterly_rows.attrs = {}
    aligned_groups: list[pd.DataFrame] = []

    # 4. Align quarterly descriptors independently for each Ticker.
    for ticker, ticker_rows in panel_rows.groupby("Ticker", sort=False):
        left = ticker_rows.sort_values("Date")
        right = quarterly_rows[
            quarterly_rows["Ticker"] == ticker
        ].sort_values("Date")

        if right.empty:
            # 7. A Ticker without quarterly history keeps explicit NaN values.
            aligned = left.copy()
            for column in expected_quarterly_columns:
                aligned[column] = np.nan
        else:
            # 5. direction="backward" prevents using a descriptor dated after
            # the target trading day; it never turns future data into history.
            # 7. Missing quarterly values remain NaN rather than being filled.
            aligned = pd.merge_asof(
                left,
                right.drop(columns="Ticker"),
                on="Date",
                direction="backward",
                allow_exact_matches=True,
            )
        aligned["Ticker"] = ticker
        aligned_groups.append(aligned)

    if aligned_groups:
        panel = pd.concat(aligned_groups, ignore_index=True)
    else:
        panel = panel_rows.copy()
        for column in expected_quarterly_columns:
            panel[column] = np.nan

    # 6. Add static GICS data by Ticker (many price rows to one industry row).
    panel = panel.merge(
        industries_for_join.reset_index(),
        on="Ticker",
        how="left",
        validate="many_to_one",
    )
    # 8. Rebuild and sort the final Date/Ticker MultiIndex.
    panel = panel.set_index(STANDARD_INDEX_NAMES).sort_index()

    # 8-9. Reject changed/duplicate price observations so the result continues
    # to satisfy the common Factor input-grid contract.
    if len(panel) != len(prices):
        raise ValueError(
            "feature joins changed the number of price observations"
        )
    if not panel.index.equals(original_index.sort_values()):
        raise ValueError("feature joins changed the price observation grid")
    if panel.index.has_duplicates:
        raise ValueError("final panel index must be unique")

    panel.attrs["point_in_time_status"] = {
        "prices": "observed",
        "daily_descriptors": "Bloomberg historical observations",
        "quarterly_descriptors": "unverified",
    }
    panel.attrs["quarterly_date_semantics"] = (
        "Bloomberg Period=Q observation date; not confirmed filing date"
    )
    return panel


def load_us_market_workbook(path: str | Path) -> pd.DataFrame:
    """Public entry point: Bloomberg workbook to standard factor panel.

    Acceptance criteria
    -------------------
    - The original workbook is never modified.
    - The workbook is not reopened once per sheet.
    - Cached Bloomberg errors become NaN plus diagnostics.
    - The returned panel can be passed to ``SizeFactor.compute``.
    """
    # 1. Validate the file path before attempting to open the workbook.
    excel_path = _validate_excel_path(path)

    # 2. Open once, read-only, with cached Bloomberg values rather than formulas.
    workbook = load_workbook(
        excel_path,
        read_only=True,
        data_only=True,
        keep_links=True,
    )
    try:
        # 1. Inspect the required workbook schema before parsing any sheet.
        missing_sheets = sorted(
            REQUIRED_US_MARKET_SHEETS - set(workbook.sheetnames)
        )
        if missing_sheets:
            raise ValueError(
                f"workbook is missing required sheets: {missing_sheets}"
            )

        # 3. Load the ordered SPX Ticker universe.
        tickers = load_universe_sheet(workbook)

        # 4. Parse wide [Date, PX_LAST] price blocks into a long table.
        prices = load_prices_sheet(workbook, tickers)

        # 5. Parse daily and quarterly descriptor blocks separately.
        daily, quarterly = load_descriptors_sheet(workbook, tickers)

        # 6. Parse the static GICS industry mapping.
        industries = load_industries_sheet(workbook, tickers)

        # 7. Assemble the standard Date/Ticker factor input panel.
        panel = build_us_market_panel(
            prices,
            daily,
            quarterly,
            industries,
        )

        panel.attrs["source_path"] = str(excel_path)
        panel.attrs["universe_size"] = len(tickers)
        panel.attrs["price_diagnostics"] = prices.attrs.get(
            "diagnostics",
            [],
        )
        panel.attrs["descriptor_diagnostics"] = daily.attrs.get(
            "diagnostics",
            [],
        )
        # 9. Return the final panel with provenance and diagnostics metadata.
        return panel
    finally:
        # 8. Always close the workbook, including when parsing raises an error.
        workbook.close()
